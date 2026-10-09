import uuid
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from contributions.models import Contribution
from contributions.services import save_contributions
from members.models import Member

from .models import ExcelUpload, ExcelUploadRow


def create_excel_upload(form, uploaded_by):
    upload = form.save(commit=False)
    upload.uploaded_by = uploaded_by
    upload.file_name = upload.file.name
    upload.upload_reference = f"UPL-{uuid.uuid4().hex[:10].upper()}"
    upload.save()
    return process_excel_upload(upload)


@transaction.atomic
def process_excel_upload(upload: ExcelUpload):
    upload.rows.all().delete()

    upload.file.open("rb")
    workbook = None
    try:
        workbook = load_workbook(upload.file, read_only=True, data_only=True)
        sheet = workbook.active

        total_rows = 0
        valid_rows = 0
        failed_rows = 0
        total_amount = Decimal("0.00")

        for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            bahasha_number = row[0] if len(row) > 0 else None
            raw_amount = row[1] if len(row) > 1 else None

            if bahasha_number is None and raw_amount is None:
                continue

            total_rows += 1
            normalized_number = str(bahasha_number or "").strip()

            try:
                amount = Decimal(str(raw_amount))
                if not amount.is_finite() or amount <= 0:
                    raise ValueError
            except (InvalidOperation, TypeError, ValueError):
                ExcelUploadRow.objects.create(
                    upload=upload,
                    row_number=index,
                    bahasha_number=normalized_number,
                    amount=Decimal("0.00"),
                    validation_status="INVALID",
                    error_message="Amount must be a positive number",
                )
                failed_rows += 1
                continue

            if not normalized_number:
                ExcelUploadRow.objects.create(
                    upload=upload,
                    row_number=index,
                    bahasha_number="",
                    amount=amount,
                    validation_status="INVALID",
                    error_message="Bahasha number is missing",
                )
                failed_rows += 1
                continue

            try:
                member = Member.objects.get(
                    bahasha_number=normalized_number,
                    church=upload.church,
                    approval_status="APPROVED",
                )
            except Member.DoesNotExist:
                ExcelUploadRow.objects.create(
                    upload=upload,
                    row_number=index,
                    bahasha_number=normalized_number,
                    amount=amount,
                    validation_status="INVALID",
                    error_message="Bahasha number not found or member not approved",
                )
                failed_rows += 1
                continue

            ExcelUploadRow.objects.create(
                upload=upload,
                row_number=index,
                bahasha_number=normalized_number,
                amount=amount,
                resolved_member=member,
                validation_status="VALID",
            )
            valid_rows += 1
            total_amount += amount
    finally:
        if workbook is not None:
            workbook.close()
        upload.file.close()

    upload.total_rows = total_rows
    upload.valid_rows = valid_rows
    upload.failed_rows = failed_rows
    upload.total_amount = total_amount
    upload.status = "VALIDATED" if failed_rows == 0 else "FAILED"
    upload.save()

    return upload

@transaction.atomic
def approve_excel_upload(upload: ExcelUpload, approved_by):
    upload = ExcelUpload.objects.select_for_update().get(pk=upload.pk)
    if upload.status != "VALIDATED":
        raise ValueError("Only validated uploads can be approved.")
    valid_rows = upload.rows.filter(validation_status="VALID")

    contributions = []
    for row in valid_rows:
        contributions.append(Contribution(
            church=upload.church,
            member=row.resolved_member,
            bahasha_number=row.bahasha_number,
            financial_year=upload.financial_year,
            contribution_week=upload.contribution_week,
            category=upload.selected_category,
            amount=row.amount,
            contribution_date=upload.contribution_week.sunday_date,
            source="EXCEL_UPLOAD",
            status="POSTED",
            reference_number=f"EXCEL-{upload.id}-{row.id}",
            posted_by=approved_by,
        ))

    save_contributions(contributions)

    upload.status = "POSTED"
    upload.approved_by = approved_by
    upload.approved_at = timezone.now()
    upload.save()

    return upload
